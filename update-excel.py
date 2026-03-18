#!/usr/bin/env python3
"""
Update Yippy-Sales-2026.xlsx with March 2026 data
"""

import openpyxl
from openpyxl.utils import get_column_letter

# File path
EXCEL_PATH = '/mnt/c/Users/Yippyops/Documents/Yippy/Operations/Financial-Tracking/Yippy-Sales-2026.xlsx'

# March 2026 data
MARCH_DATA = {
    'revenue': {
        'For the Course (5 Pack)': 4593.85,
        'For the Desk (5 Pack)': 802.80,
        'Yippy Value Pack (10 Pack)': 1484.79,
        'TOTAL': 6881.44
    },
    'units': {
        'For the Course (5 Pack)': 115,
        'For the Desk (5 Pack)': 20,
        'Yippy Value Pack (10 Pack)': 21,
        'TOTAL': 156
    },
    'cogs': 3438.48
}

def main():
    print('📂 Opening Excel file...')
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    # Tab 1: Sales 2026
    print('📝 Updating Sales 2026 tab...')
    sales_sheet = wb['Sales 2026']
    
    # Find March column (should be column D if Jan=B, Feb=C)
    # Find the header row first
    march_col = None
    for col in range(1, 20):
        cell_value = sales_sheet.cell(row=1, column=col).value
        if cell_value == 'Mar 2026' or cell_value == 'March 2026':
            march_col = col
            break
    
    if not march_col:
        # March column doesn't exist, create it
        # Assuming Jan is col B (2), Feb is col C (3), March should be D (4)
        march_col = 4
        sales_sheet.cell(row=1, column=march_col, value='Mar 2026')
        print(f'   ✅ Created March column at {get_column_letter(march_col)}')
    else:
        print(f'   ✅ Found March column at {get_column_letter(march_col)}')
    
    # Update revenue section (starting around row 3)
    # Look for product names in column A
    for row in range(2, 100):
        cell_value = sales_sheet.cell(row=row, column=1).value
        if cell_value in MARCH_DATA['revenue']:
            sales_sheet.cell(row=row, column=march_col, value=MARCH_DATA['revenue'][cell_value])
            print(f'   💰 {cell_value}: ${MARCH_DATA["revenue"][cell_value]:.2f}')
    
    # Update units section (look for "Units Sold" section)
    units_section_found = False
    for row in range(2, 100):
        cell_value = sales_sheet.cell(row=row, column=1).value
        if cell_value == 'Units Sold' or cell_value == 'UNITS SOLD':
            units_section_found = True
            units_start_row = row + 1
            break
    
    if units_section_found:
        for row in range(units_start_row, units_start_row + 20):
            cell_value = sales_sheet.cell(row=row, column=1).value
            if cell_value in MARCH_DATA['units']:
                sales_sheet.cell(row=row, column=march_col, value=MARCH_DATA['units'][cell_value])
                print(f'   📦 {cell_value}: {MARCH_DATA["units"][cell_value]} units')
    
    # Tab 2: Expenses
    print('\n💵 Updating Expenses tab...')
    expenses_sheet = wb['Expenses']
    
    # Find COGS row and March column
    march_col_expenses = None
    cogs_row = None
    
    # Find March column header
    for col in range(1, 20):
        cell_value = expenses_sheet.cell(row=1, column=col).value
        if cell_value == 'Mar 2026' or cell_value == 'March 2026':
            march_col_expenses = col
            break
    
    if not march_col_expenses:
        march_col_expenses = 4  # Same as sales sheet
        expenses_sheet.cell(row=1, column=march_col_expenses, value='Mar 2026')
        print(f'   ✅ Created March column at {get_column_letter(march_col_expenses)}')
    
    # Find COGS row
    for row in range(2, 100):
        cell_value = expenses_sheet.cell(row=row, column=1).value
        if cell_value and ('COGS' in str(cell_value) or 'Cost of Goods' in str(cell_value)):
            cogs_row = row
            break
    
    if cogs_row:
        # Check if it's a formula row - if so, don't overwrite (let it auto-calculate)
        cell = expenses_sheet.cell(row=cogs_row, column=march_col_expenses)
        if cell.value is None or not isinstance(cell.value, str) or not cell.value.startswith('='):
            expenses_sheet.cell(row=cogs_row, column=march_col_expenses, value=MARCH_DATA['cogs'])
            print(f'   ✅ COGS updated: ${MARCH_DATA["cogs"]:.2f}')
        else:
            print(f'   ℹ️  COGS is a formula - will auto-calculate')
    
    # Save
    print('\n💾 Saving changes...')
    wb.save(EXCEL_PATH)
    print('✅ Excel file updated successfully!')
    
    print('\n━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━')
    print('📊 MARCH 2026 SUMMARY:')
    print('━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━')
    print(f'Revenue:      ${MARCH_DATA["revenue"]["TOTAL"]:.2f}')
    print(f'Units:        {MARCH_DATA["units"]["TOTAL"]}')
    print(f'COGS:         ${MARCH_DATA["cogs"]:.2f}')
    print(f'Gross Profit: ${MARCH_DATA["revenue"]["TOTAL"] - MARCH_DATA["cogs"]:.2f}')
    print(f'Gross Margin: {(MARCH_DATA["revenue"]["TOTAL"] - MARCH_DATA["cogs"]) / MARCH_DATA["revenue"]["TOTAL"] * 100:.1f}%')

if __name__ == '__main__':
    main()
